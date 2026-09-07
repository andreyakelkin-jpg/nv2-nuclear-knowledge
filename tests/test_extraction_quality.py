from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path


PLUGIN_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PLUGIN_ROOT / "scripts"))

from extraction_quality import extract_text, extraction_quality  # noqa: E402


class ExtractionQualityTests(unittest.TestCase):
    def test_page_quality_is_evaluated_per_marked_page(self) -> None:
        text = (
            "===== СТРАНИЦА 1 =====\n" + "полный текст " * 12
            + "\n===== СТРАНИЦА 2 =====\n"
        )
        quality = extraction_quality(text, "pypdf")
        self.assertEqual(2, quality["pages_detected"])
        self.assertEqual([2], quality["empty_or_sparse_pages"])
        self.assertTrue(quality["requires_visual_review"])

    def test_pdf_without_markers_is_flagged_even_when_text_is_long(self) -> None:
        quality = extraction_quality("текст " * 100, "pypdf")
        self.assertEqual(0, quality["pages_detected"])
        self.assertIn("pdf_page_markers_missing", quality["warnings"])
        self.assertTrue(quality["requires_visual_review"])

    def test_plain_text_extracts_without_pdf_page_warning(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.txt"
            path.write_text("Нормативный текст", encoding="utf-8")
            text, method = extract_text(path)
        self.assertEqual("plain-text", method)
        self.assertEqual([], extraction_quality(text, method)["warnings"])

    def test_sparse_pypdf_result_uses_text_extraction_fallback(self) -> None:
        with patch("extraction_quality._extract_with_pypdf", return_value=["x", "y"]), \
             patch("extraction_quality._extract_with_pdfplumber", return_value=["достаточный текст"]):
            text, method = extract_text(Path("fixture.pdf"))
        self.assertEqual("pdfplumber", method)
        self.assertIn("===== СТРАНИЦА 1 =====", text)
        self.assertIn("достаточный текст", text)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx not installed")
    def test_docx_keeps_paragraph_table_paragraph_order(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.docx"
            document = Document()
            document.add_paragraph("Первый абзац")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Ячейка A"
            table.cell(0, 1).text = "Ячейка B"
            document.add_paragraph("Последний абзац")
            document.save(path)
            text, method = extract_text(path)
        self.assertEqual("python-docx", method)
        self.assertLess(text.index("Первый абзац"), text.index("[ТАБЛИЦА 1]"))
        self.assertLess(text.index("[ТАБЛИЦА 1]"), text.index("Последний абзац"))
        self.assertIn("Ячейка A | Ячейка B", text)


if __name__ == "__main__":
    unittest.main()
